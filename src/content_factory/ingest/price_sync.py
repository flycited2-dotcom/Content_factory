"""Mirror a verified supplier release into the factory's independent price slots.

This updates files only: no task queue, photoagent or publication operations.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import tempfile

SOURCE = 'telegram-nikita'
SLOT = 'manual__быттехопт_auto'


def _json(path: Path) -> dict:
    try:
        value = json.loads(path.read_text(encoding='utf-8'))
        return value if isinstance(value, dict) else {}
    except (OSError, ValueError):
        return {}


def _write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.tmp')
    with temp.open('w', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temp, path)


def _rows(catalog: Path) -> tuple[dict, list[tuple], int]:
    from content_factory.ready_price import _source_problem
    with sqlite3.connect(catalog.resolve().as_uri() + '?mode=ro', uri=True) as db:
        db.execute('BEGIN')
        problem = _source_problem(db)
        if problem:
            raise ValueError(problem)
        digest, generated, raw = db.execute(
            "SELECT sha256,generated_at,manifest FROM releases WHERE source=? "
            "AND status='accepted' ORDER BY generated_at DESC LIMIT 1", (SOURCE,)).fetchone()
        manifest = json.loads(raw)
        rows = db.execute('SELECT article,data FROM catalog WHERE source=? AND sha256=? AND present=1',
                          (SOURCE, digest)).fetchall()
    if not str(manifest.get('original_filename', '')).startswith('БытТехОпт_'):
        raise ValueError('unexpected_supplier_release')
    output, excluded = [], 0
    for article, raw in rows:
        data = json.loads(raw)
        if data.get('availability') != 'supplier_price_present':
            excluded += 1
            continue
        # source_price is the ready price before Avito's extra markup.
        try:
            price = Decimal(str(data['source_price']))
        except (KeyError, InvalidOperation):
            raise ValueError('invalid_source_price') from None
        if not price.is_finite() or price <= 0 or not data.get('name') or not data.get('brand'):
            raise ValueError('invalid_supplier_row')
        output.append((str(data.get('group') or 'Другие товары'), str(article),
                       str(data['brand']), str(data['name']),
                       int(price.quantize(Decimal('1'), rounding=ROUND_HALF_UP))))
    if not output:
        raise ValueError('empty_verified_supplier_price')
    return {'release_sha256': digest, 'filename': manifest['original_filename'],
            'issue_date': manifest.get('issue_date', ''), 'generated_at': generated,
            'supplier_updated_at': manifest['provenance']['supplier']['modified_at']}, output, excluded


def _adopt_preferences(prices: Path) -> None:
    """Migrate the confirmed old BT slot's preferences once, preserving owner edits."""
    old = 'manual__быттехопт_20260704'
    markup_path = prices / 'markups.json'
    markups = _json(markup_path)
    if old in markups and SLOT not in markups:
        markups[SLOT] = markups[old]
        _write_json(markup_path, markups)
    toggle_path = prices / 'telegram_sources.json'
    toggles = _json(toggle_path)
    disabled = list(toggles.get('disabled') or [])
    if old in disabled and SLOT not in disabled:
        toggles['disabled'] = disabled + [SLOT]
        _write_json(toggle_path, toggles)


def _sync_prices_unlocked(catalog: Path, prices: Path) -> dict:
    import openpyxl
    metadata, rows, excluded = _rows(catalog)
    prices.mkdir(parents=True, exist_ok=True)
    target = prices / f'{SLOT}.xlsx'
    registry_path = prices / 'source_refresh.json'
    registry = _json(registry_path)
    previous = registry.get(SLOT, {})
    if (previous.get('release_sha256') == metadata['release_sha256'] and target.exists()
            and hashlib.sha256(target.read_bytes()).hexdigest() == previous.get('sha256')):
        return {'status': 'unchanged', 'slot': SLOT, 'items': len(rows), 'excluded': excluded}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Прайс'
    ws.append(['№', 'Артикул', 'Бренд', 'Наименование', 'Цена (руб.)', 'Заказ (шт.)'])
    last_group = None
    for number, (group, article, brand, name, price) in enumerate(sorted(rows), 1):
        if group != last_group:
            ws.append([group])
            ws.cell(ws.max_row, 1).data_type = 's'
            last_group = group
        # Force supplier strings to text; an XLSX is data, never executable formulas.
        ws.append([number, article, brand, name, price, None])
        for col in (2, 3, 4):
            ws.cell(ws.max_row, col).data_type = 's'
    fd, temp_name = tempfile.mkstemp(prefix='.supplier-price-', suffix='.xlsx', dir=prices)
    os.close(fd)
    temp = Path(temp_name)
    try:
        wb.save(temp)
        wb.close()
        from content_factory.ingest.excel_price import parse_price_xlsx
        if len(parse_price_xlsx(temp)) != len(rows):
            raise ValueError('generated_price_validation_failed')
        digest = hashlib.sha256(temp.read_bytes()).hexdigest()
        if not target.exists():
            _adopt_preferences(prices)
        os.replace(temp, target)
        registry[SLOT] = {**metadata, 'display_name': 'БытТехОпт (авто)',
                          'updated_at': datetime.now(timezone.utc).isoformat(),
                          'sha256': digest, 'items': len(rows), 'excluded': excluded}
        _write_json(registry_path, registry)
    finally:
        temp.unlink(missing_ok=True)
    return {'status': 'updated', 'slot': SLOT, 'items': len(rows), 'excluded': excluded,
            'release_sha256': metadata['release_sha256']}


def sync_prices(catalog: Path, prices: Path) -> dict:
    # Serialize timer/manual recovery before reading the latest release, so an
    # older invocation cannot finish after and overwrite a newer supplier file.
    prices.mkdir(parents=True, exist_ok=True)
    lock = sqlite3.connect(prices / '.price-sync-lock.sqlite', timeout=60)
    try:
        lock.execute('BEGIN IMMEDIATE')
        return _sync_prices_unlocked(catalog, prices)
    finally:
        lock.rollback()
        lock.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--catalog', type=Path, default=Path('/opt/avito-bridge/state/ready-price/catalog.sqlite'))
    parser.add_argument('--prices', type=Path)
    args = parser.parse_args()
    if args.prices is None:
        from content_factory.config import load_config
        args.prices = Path(load_config('config/config.yaml').state.db).parent / 'prices'
    report = args.prices / 'price-sync-status.json'
    try:
        result = sync_prices(args.catalog, args.prices)
    except Exception as exc:
        result = {'status': 'deferred', 'reason': str(exc) if isinstance(exc, ValueError) else type(exc).__name__}
    result['checked_at'] = datetime.now(timezone.utc).isoformat()
    _write_json(report, result)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
